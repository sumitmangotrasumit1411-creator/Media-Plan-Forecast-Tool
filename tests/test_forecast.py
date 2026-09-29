import io
from decimal import Decimal
import pandas as pd
import pytest
from openpyxl import load_workbook
from forecast import (DEFAULT_CHANNEL_SPLIT, exact_channel_split, rebalance_channel_split,
                      run_forecast, monthly_forecast)
from exporter import build_excel_media_plan

CHANNELS = list(DEFAULT_CHANNEL_SPLIT)

def forecast(**kw):
    return run_forecast(1000000, 100000, 400000, 20, **kw)

def test_exact():
    assert list(exact_channel_split(70,20,10).values()) == [.7,.2,.1]
    for values in [(70,20,20), (30,40,30), (70,15,15)]:
        with pytest.raises(ValueError):
            exact_channel_split(*values)

@pytest.mark.parametrize('channel', CHANNELS)
def test_all_slider_values(channel):
    current = DEFAULT_CHANNEL_SPLIT
    for pct in range(1001):
        current = rebalance_channel_split(channel, pct/10, current)
        sp,sb,sd = current.values()
        assert sp > sb > sd >= 0
        assert sum(Decimal(str(v)) for v in current.values()) == 1

def test_sp_sixty():
    result = rebalance_channel_split(CHANNELS[0],60)
    assert list(result.values()) == [.6,.286,.114]

@pytest.mark.parametrize('key,value', [('override_target_revenue',1500000),('override_ad_spend',150000),('override_ad_sales',600000),('override_roas',6),('override_tacos',15)])
def test_targets(key,value):
    a,b = forecast(**{key:value}),forecast(**{key:value*1.2})
    assert any(a[k] != b[k] for k in ['recommended_spend','target_ad_sales','target_revenue'])
    for s in [a,b]:
        df,_ = monthly_forecast(None,s['growth_pct'],1000000,annual_spend_override=s['recommended_spend'],annual_sales_override=s['target_ad_sales'])
        assert round(df['Projected Spend ($)'].sum(),2) == s['recommended_spend']
        assert round(df['Projected Ad Sales ($)'].sum(),2) == s['target_ad_sales']
        assert round(sum(x['budget'] for x in s['channel_allocation'].values()),2) == s['recommended_spend']
        for col, alloc in zip(['SP Budget ($)','SB Budget ($)','SD Budget ($)'], s['channel_allocation'].values()):
            assert round(df[col].sum(),2) == alloc['budget']
        assert ((df[['SP Budget ($)','SB Budget ($)','SD Budget ($)']].sum(axis=1)-df['Projected Spend ($)']).abs()<1e-8).all()

def test_precedence_and_mix():
    s=forecast(override_ad_spend=123456,override_roas=7,override_tacos=9,override_ad_sales=600000,override_target_revenue=1500000)
    assert s['recommended_spend']==123456
    assert s['target_ad_sales']==600000
    assert s['target_revenue']==1500000
    assert forecast(override_roas=5,override_tacos=9)['projected_roas']==5
    assert forecast()['recommended_spend'] != forecast(custom_channel_split=exact_channel_split(70,20,10))['recommended_spend']

def test_zero_override():
    trend=pd.DataFrame({'month':[1], 'spend':[100], 'ad_sales':[200]})
    df,_=monthly_forecast(trend,10,1000,annual_spend_override=0,annual_sales_override=0)
    assert df['Projected Spend ($)'].sum()==0
    assert df['Projected Ad Sales ($)'].sum()==0

def test_export_custom():
    s=forecast(override_ad_spend=123456)
    df,_=monthly_forecast(None,20,1000000,annual_spend_override=s['recommended_spend'],annual_sales_override=s['target_ad_sales'])
    data=build_excel_media_plan({}, {}, [s],pd.DataFrame(),pd.DataFrame(),monthly_df=df)
    wb=load_workbook(io.BytesIO(data))
    ws=wb['Scenarios']
    assert ws['A2'].value=='Custom'
    assert ws['F2'].value==123456
    assert '$' in ws['J2'].number_format
    monthly=wb['Monthly Media Plan']
    col=next(c.column for c in monthly[1] if c.value=='Projected ROAS')
    assert '$' in monthly.cell(2,col).number_format
